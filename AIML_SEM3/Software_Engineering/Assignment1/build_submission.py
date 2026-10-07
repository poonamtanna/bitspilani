"""Generate the Group 54 assignment notebook and Word report."""

from pathlib import Path
import base64
import re
import subprocess
import sys
import textwrap
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import matplotlib.pyplot as plt
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
import nbformat


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "submission_assets"
ASSETS.mkdir(exist_ok=True)
MEMBER_HEADERS = [
    "Sl. No", "BITS ID", "Name", "Contribution (Qualitative)", "Percentage (out of 100)",
]
MEMBER_ROWS = [
    ["1", "2025AE05017", "POONAM TANNA", "Contributed across most project areas, collaborating on problem formulation, GR4ML views, system architecture, pattern implementation, Streamlit UI, and report generation.", "40"],
    ["2", "2025AE05791", "Mohammed Asif S", "Requirements formulation: domain and problem statement, users and stakeholders, GR4ML measurable goals and indicators table, functional and non-functional requirements with acceptance criteria, and justification of the top three quality requirements.", "20"],
    ["3", "2025AF05145", "NAVEEN.K.B", "GR4ML modelling: Business View, Analytics Design View and Data Preparation View (tables and diagrams), cross-view traceability, softgoal contribution links and decomposition correctness.", "20"],
    ["4", "2025AF05136", "SIMRAN SURI", "Evaluation and reporting: model evaluation experiments (threshold sweep, fairness audit by subgroup), application testing and screenshots with captions, notebook execution and saved outputs, report compilation, formatting and consistency review between docx and notebook.", "20"],
]


def member_markdown_table():
    lines = ["| " + " | ".join(MEMBER_HEADERS) + " |", "|" + "|".join(["---"] * len(MEMBER_HEADERS)) + "|"]
    lines.extend("| " + " | ".join(row) + " |" for row in MEMBER_ROWS)
    return "\n".join(lines)


def markdown_table(headers, rows):
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    lines.extend("| " + " | ".join(str(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def image_markdown_cell(filename, description):
    encoded = base64.b64encode((ASSETS / filename).read_bytes()).decode("ascii")
    cell = nbformat.v4.new_markdown_cell(f"![{description}](attachment:{filename})")
    cell.attachments = {filename: {"image/png": encoded}}
    return cell


def save_figure(figure, output_path, dpi):
    output_path = Path(output_path)
    temporary_path = output_path.with_name(f".{output_path.stem}-{uuid4().hex}.png")
    figure.savefig(temporary_path, dpi=dpi, bbox_inches="tight")
    temporary_path.replace(output_path)
GQI_HEADERS = [
    "Goal ID", "GR4ML Goal Type", "Goal Statement", "Indicator (KPI)",
    "Target Threshold", "Prototype Holdout Result (t=0.50)",
]
def collect_submission_evidence():
    from credit_risk.fairness import audit_recall_by_group
    from credit_risk.service import evaluate_model, fit_model
    from credit_risk.strategies import STRATEGIES, get_strategy

    thresholds = [round(0.30 + 0.05 * index, 2) for index in range(7)]
    models = {}
    baseline = {}
    sweep_rows = []
    fairness = {}
    fairness_rows = []
    for strategy_name in STRATEGIES:
        model = fit_model(get_strategy(strategy_name))
        models[strategy_name] = model
        fairness[strategy_name] = audit_recall_by_group(model, threshold=0.5)
        for dimension, groups in fairness[strategy_name]["recall_by_group"].items():
            fairness_rows.extend(
                {"strategy": strategy_name, "dimension": dimension, "group": group,
                 "recall": recall}
                for group, recall in groups.items()
            )
        for threshold in thresholds:
            metrics = evaluate_model(model, threshold)
            row = {"strategy": strategy_name, "threshold": threshold, **metrics}
            sweep_rows.append(row)
            if threshold == 0.5:
                baseline[strategy_name] = metrics

    passing_thresholds = [
        row for row in sweep_rows
        if row["recall"] >= 0.70 and row["specificity"] >= 0.70
    ]
    return {
        "baseline": baseline,
        "sweep": sweep_rows,
        "fairness": fairness,
        "fairness_rows": fairness_rows,
        "threshold_target_met": bool(passing_thresholds),
        "passing_thresholds": passing_thresholds,
    }


def build_gqi_rows(evidence):
    baseline = evidence["baseline"]
    fairness = evidence["fairness"]
    lr = baseline["Logistic regression"]
    rf = baseline["Random forest"]
    fairness_result = "; ".join(
        f"{name}: {data['max_gap_pp']:.2f}pp ({'meets' if data['meets_target'] else 'does not meet'} <=5pp)"
        for name, data in fairness.items()
    )
    return [
        ["G-B1", "Strategic Goal", "Support timely review of accounts with elevated default risk.",
         "Default recall on the positive class", ">=70% on independent later-period validation",
         f"Logistic regression {lr['recall']:.2%}; Random forest {rf['recall']:.2%}."],
        ["DG-1", "Decision Goal", "Decide account review priority for analyst follow-up.",
         "Default recall and specificity", ">=70% for each on independent validation",
         f"Recall {lr['recall']:.2%}/{rf['recall']:.2%}; specificity {lr['specificity']:.2%}/{rf['specificity']:.2%} (LR/RF)."],
        ["QG-1", "Question Goal", "What is the account's next-month default probability?",
         "Probability output in [0, 1] and correct threshold flag", "Valid probability and correct threshold behavior for every request",
         "Probability and threshold boundary are covered by the automated test suite; no calibration claim."],
        ["IND-1", "Indicator", "Measure default detection effectiveness.",
         "Recall = TP / (TP + FN)", ">=70% on independent later-period validation",
         f"Logistic regression {lr['recall']:.2%}; Random forest {rf['recall']:.2%}."],
        ["IND-2", "Indicator", "Measure correct non-default identification.",
         "Specificity = TN / (TN + FP)", ">=70% on independent later-period validation",
         f"Logistic regression {lr['specificity']:.2%}; Random forest {rf['specificity']:.2%}."],
        ["Q-B2", "Softgoal", "Ensure no material demographic disparity in recall.",
         "Maximum recall gap by sex or age band", "<=5 percentage points before pilot",
         fairness_result + ". Coursework audit only; not a governed audit."],
    ]


def run_test_suite():
    result = subprocess.run(
        [sys.executable, "-m", "pytest", "tests/test_credit_risk.py", "-q"],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )
    output = (result.stdout + result.stderr).strip()
    match = re.search(r"(\d+) passed", output)
    if result.returncode != 0 or not match:
        raise RuntimeError(f"pytest failed; report was not generated:\n{output}")
    return int(match.group(1)), output
SOURCE_FILES = [
    ROOT / "credit_risk" / "data.py",
    ROOT / "credit_risk" / "strategies.py",
    ROOT / "credit_risk" / "service.py",
    ROOT / "credit_risk" / "fairness.py",
    ROOT / "app.py",
    ROOT / "api.py",
    ROOT / "tests" / "test_credit_risk.py",
]


def draw_architecture():
    figure, axis = plt.subplots(figsize=(15, 8))
    figure.patch.set_facecolor("#f4f6f3")
    axis.set_facecolor("#f4f6f3")
    axis.set_xlim(0, 16)
    axis.set_ylim(0, 9.5)
    axis.axis("off")
    nodes = {
        "ui": (1.8, 7.0), "api": (5.8, 7.0), "service": (3.8, 4.6),
        "data": (1.8, 2.0), "cache": (5.8, 2.0), "governance": (3.8, 0.8),
        "training": (9.8, 7.0), "strategies": (13.5, 7.0),
        "artifact": (13.5, 4.1), "scoring": (9.8, 4.1),
    }

    def component(key, title, body, color, width=2.8):
        x, y = nodes[key]
        axis.add_patch(FancyBboxPatch(
            (x - width / 2, y - 0.55), width, 1.1,
            boxstyle="round,pad=0.05,rounding_size=0.1", linewidth=1.4,
            edgecolor="#29423a", facecolor=color, zorder=2,
        ))
        axis.text(x, y + 0.2, title, ha="center", va="center", fontsize=9,
                  weight="bold", color="#18332b", zorder=3)
        axis.text(x, y - 0.2, body, ha="center", va="center", fontsize=7.2,
                  color="#29423a", zorder=3)

    axis.add_patch(FancyBboxPatch(
        (0.35, 0.25), 7.1, 7.65, boxstyle="round,pad=0.04,rounding_size=0.1",
        facecolor="#f8faf8", edgecolor="#55756a", linewidth=1.4, zorder=0,
    ))
    axis.add_patch(FancyBboxPatch(
        (7.75, 2.7), 7.85, 5.2, boxstyle="round,pad=0.04,rounding_size=0.1",
        facecolor="#fffaf0", edgecolor="#b28c4b", linewidth=1.4, zorder=0,
    ))
    axis.text(3.9, 8.08, "NON-ML COMPONENTS", ha="center", fontsize=10,
              weight="bold", color="#29423a")
    axis.text(11.7, 8.08, "ML COMPONENTS", ha="center", fontsize=10,
              weight="bold", color="#71501d")

    edges = [
        ("ui", "service", "account inputs", 0.0, 0.65, 0.25),
        ("service", "ui", "probability + review flag", -0.2, -0.65, -0.25),
        ("api", "service", "API request / response", 0.0, 0.7, 0.05),
        ("service", "scoring", "features + threshold", 0.0, 0.0, 0.2),
        ("service", "training", "cache miss / retrain", 0.0, 0.2, 0.35),
        ("data", "training", "stratified data", 0.0, 0.0, -0.25),
        ("training", "strategies", "fit selected model", 0.0, 0.0, 0.2),
        ("strategies", "artifact", "fitted model", 0.0, 0.0, 0.2),
        ("artifact", "scoring", "positive-class score", 0.0, 0.0, 0.2),
        ("cache", "scoring", "cache hit", 0.0, 0.0, 0.2),
        ("service", "governance", "aggregate metrics only", 0.0, 0.5, 0.0),
    ]
    for source, target, label, bend, label_dx, label_dy in edges:
        start, end = nodes[source], nodes[target]
        axis.add_patch(FancyArrowPatch(
            start, end, arrowstyle="-|>", mutation_scale=13, linewidth=1.15,
            color="#45685a", connectionstyle=f"arc3,rad={bend}",
            shrinkA=50, shrinkB=50, zorder=1,
        ))
        axis.text(
            (start[0] + end[0]) / 2 + label_dx,
            (start[1] + end[1]) / 2 + label_dy,
            label, ha="center", va="center", fontsize=6.5, color="#344d43",
            bbox={"facecolor": "#f4f6f3", "edgecolor": "none", "pad": 1}, zorder=4,
        )

    component("ui", "Streamlit UI", "controls + illustrative inputs", "#dce9df")
    component("api", "FastAPI Client/Server", "POST /predict | GET /metrics", "#e8eef5", 3.0)
    component("service", "Application Service", "shared fit, evaluate, predict", "#dce9df", 3.0)
    component("data", "Data Adapter", "bundled CSV; 19 model features", "#e8eef5")
    component("cache", "Cache / Registry", "process-local fitted model", "#e8eef5")
    component("governance", "Governance Logger", "aggregate metrics only", "#f2e6e6", 3.0)
    component("training", "Training / Evaluation", "stratified holdout + KPIs", "#f9f1dc", 3.0)
    component("strategies", "Model Strategies", "logistic regression / forest", "#f4e5c5", 3.0)
    component("artifact", "Fitted-model artifact", "positive-class probability model", "#f4e5c5", 3.0)
    component("scoring", "Scoring", "default probability + review flag", "#f9f1dc", 3.0)
    axis.text(8, 0.08, "No identifiers, account-level persistence, or automated credit action",
              ha="center", fontsize=8.5, color="#7c3f32", weight="bold")
    axis.set_title("Group 54 | Credit Default Risk Review Architecture", fontsize=15,
                   weight="bold", color="#18332b", pad=15)
    save_figure(figure, ASSETS / "architecture.png", dpi=220)
    plt.close(figure)


def draw_gr4ml_view(filename, title, nodes, edges, annotations=()):
    figure, axis = plt.subplots(figsize=(13, 8.5))
    figure.patch.set_facecolor("#fbfcfa")
    axis.set_facecolor("#fbfcfa")
    axis.set_xlim(0, 14)
    axis.set_ylim(0, 10)
    axis.axis("off")
    positions = {node["id"]: (node["x"], node["y"]) for node in nodes}

    for edge in edges:
        source, target, label, style, *routing = edge
        bend, label_dx, label_dy, shrink_a, shrink_b = (
            routing + [0.0, 0.0, 0.12, 35, 35]
        )[:5]
        x1, y1 = positions[source]
        x2, y2 = positions[target]
        axis.add_patch(FancyArrowPatch(
            (x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=12,
            linewidth=1.0, color="#50675f", linestyle="--" if style == "rule" else "-",
            shrinkA=shrink_a, shrinkB=shrink_b,
            connectionstyle=f"arc3,rad={bend}", zorder=1,
        ))
        axis.text((x1 + x2) / 2 + label_dx, (y1 + y2) / 2 + label_dy, label,
                  fontsize=7, color="#344d43", ha="center", va="center",
                  bbox={"facecolor": "#fbfcfa", "edgecolor": "none", "pad": 1.1}, zorder=3)

    colors = {
        "actor": "#f4e5c5",
        "goal": "#dce9df",
        "strategic_goal": "#dce9df",
        "decision_goal": "#dce9df",
        "question_goal": "#dce9df",
        "analytics_goal": "#dce9df",
        "softgoal": "#e8eef5",
        "task": "#f9f1dc",
        "algorithm": "#f4e5c5",
        "indicator": "#e3f2f4",
        "data_store": "#e5eee9",
        "resource": "#f2e6e6",
    }
    for node in nodes:
        x, y, node_id, label, kind = (
            node["x"], node["y"], node["id"], node["label"], node["kind"]
        )
        color = colors[kind]
        if kind == "actor":
            patch = Circle((x, y), radius=0.68, facecolor=color, edgecolor="#29423a", linewidth=1.25)
        elif kind in {"goal", "strategic_goal", "decision_goal", "question_goal", "analytics_goal"}:
            patch = Ellipse((x, y), width=2.45, height=0.92, facecolor=color, edgecolor="#29423a", linewidth=1.25)
        elif kind == "softgoal":
            patch = FancyBboxPatch(
                (x - 1.16, y - 0.43), 2.32, 0.86,
                boxstyle="round,pad=0.04,rounding_size=0.12", facecolor=color,
                edgecolor="#29423a", linewidth=1.2, linestyle="--",
            )
        elif kind in {"task", "algorithm"}:
            patch = Polygon(
                [(x - 1.15, y - 0.5), (x + 1.15, y - 0.5),
                 (x + 1.35, y), (x + 1.15, y + 0.5),
                 (x - 1.15, y + 0.5), (x - 1.35, y)],
                closed=True, facecolor=color, edgecolor="#29423a", linewidth=1.25,
            )
        elif kind == "indicator":
            patch = Polygon(
                [(x - 1.2, y - 0.42), (x + 0.9, y - 0.42),
                 (x + 1.25, y), (x + 0.9, y + 0.42), (x - 1.2, y + 0.42)],
                closed=True, facecolor=color, edgecolor="#29423a", linewidth=1.25,
            )
        elif kind == "data_store":
            width = node.get("width", 2.5)
            height = node.get("height", 0.95)
            body = Rectangle((x - width / 2, y - height / 2 + 0.12), width, height - 0.24,
                             facecolor=color, edgecolor="#29423a", linewidth=1.25)
            top = Ellipse((x, y + height / 2 - 0.12), width, 0.3,
                          facecolor=color, edgecolor="#29423a", linewidth=1.25)
            bottom = Ellipse((x, y - height / 2 + 0.12), width, 0.3,
                             facecolor=color, edgecolor="#29423a", linewidth=1.25)
            axis.add_patch(bottom)
            axis.add_patch(body)
            patch = top
        else:
            patch = Rectangle((x - 1.16, y - 0.43), 2.32, 0.86,
                              facecolor=color, edgecolor="#29423a", linewidth=1.25)
        patch.set_zorder(2)
        axis.add_patch(patch)
        axis.text(x, y + 0.12, node_id, ha="center", va="center",
                  fontsize=8.2, weight="bold", color="#18332b", zorder=4)
        axis.text(x, y - 0.15, label, ha="center", va="center",
                  fontsize=6.5, color="#29423a", zorder=4, wrap=True)

    for x, y, annotation in annotations:
        axis.text(x, y, annotation, ha="center", va="center", fontsize=7.4,
                  color="#344d43", zorder=4)

    legend_y = 0.73
    legend_items = [
        (0.65, "actor", "Actor"), (2.75, "goal", "Goal"),
        (4.8, "softgoal", "Softgoal"), (6.8, "task", "Task / algorithm"),
        (9.25, "indicator", "Indicator"), (11.45, "data_store", "Data source/store"),
    ]
    for x, kind, label in legend_items:
        color = colors[kind]
        if kind == "actor":
            marker = Circle((x, legend_y), 0.14, facecolor=color, edgecolor="#29423a", linewidth=0.8)
        elif kind in {"goal", "strategic_goal", "decision_goal", "question_goal", "analytics_goal"}:
            marker = Ellipse((x, legend_y), 0.48, 0.25, facecolor=color, edgecolor="#29423a", linewidth=0.8)
        elif kind == "softgoal":
            marker = FancyBboxPatch((x - 0.23, legend_y - 0.12), 0.46, 0.24,
                                    boxstyle="round,pad=0.01", facecolor=color,
                                    edgecolor="#29423a", linewidth=0.8, linestyle="--")
        elif kind in {"task", "algorithm"}:
            marker = Polygon([(x - 0.22, legend_y - 0.12), (x + 0.22, legend_y - 0.12),
                              (x + 0.27, legend_y), (x + 0.22, legend_y + 0.12),
                              (x - 0.22, legend_y + 0.12), (x - 0.27, legend_y)],
                             closed=True, facecolor=color, edgecolor="#29423a", linewidth=0.8)
        elif kind == "indicator":
            marker = Polygon([(x - 0.23, legend_y - 0.12), (x + 0.18, legend_y - 0.12),
                              (x + 0.27, legend_y), (x + 0.18, legend_y + 0.12),
                              (x - 0.23, legend_y + 0.12)],
                             closed=True, facecolor=color, edgecolor="#29423a", linewidth=0.8)
        elif kind == "data_store":
            marker = FancyBboxPatch((x - 0.23, legend_y - 0.12), 0.46, 0.24,
                                    boxstyle="round,pad=0.02,rounding_size=0.08",
                                    facecolor=color, edgecolor="#29423a", linewidth=0.8)
        else:
            marker = Rectangle((x - 0.23, legend_y - 0.12), 0.46, 0.24,
                               facecolor=color, edgecolor="#29423a", linewidth=0.8)
        axis.add_patch(marker)
        axis.text(x + 0.34, legend_y, label, fontsize=7.2, va="center", color="#29423a")
    axis.text(7, 0.23, "Links: AND / OR = decomposition   |   depends-on = actor dependency   |   + / ++ / - = softgoal contribution",
              ha="center", va="center", fontsize=8, color="#344d43")
    axis.set_title(title, fontsize=15, weight="bold", color="#18332b", pad=16)
    save_figure(figure, ASSETS / filename, dpi=240)
    plt.close(figure)


def draw_gr4ml_business_view():
    draw_gr4ml_view("gr4ml_business_view.png", "GR4ML Business View | Credit Default Risk Review", [
        {"id": "A-1", "label": "Credit-risk analyst", "kind": "actor", "x": 1.8, "y": 8.2},
        {"id": "A-2", "label": "ML service", "kind": "actor", "x": 6.4, "y": 8.2},
        {"id": "A-3", "label": "Data steward", "kind": "actor", "x": 12.3, "y": 8.2},
        {"id": "G-B1", "label": "Timely risk\nreview", "kind": "strategic_goal", "x": 2.2, "y": 6.0},
        {"id": "G-B2", "label": "Prioritize accounts", "kind": "decision_goal", "x": 6.3, "y": 6.0},
        {"id": "G-B2a", "label": "Valid feature input", "kind": "goal", "x": 4.3, "y": 4.0},
        {"id": "G-B2b", "label": "Risk estimate + flag", "kind": "goal", "x": 8.3, "y": 4.0},
        {"id": "IND-1", "label": "Default recall", "kind": "indicator", "x": 1.8, "y": 3.0},
        {"id": "IND-2", "label": "Specificity", "kind": "indicator", "x": 4.7, "y": 3.0},
        {"id": "G-A1", "label": "Analytics goal\n(cross-view)", "kind": "analytics_goal", "x": 9.2, "y": 2.0},
        {"id": "T-B1", "label": "Investigate flagged\naccount", "kind": "task", "x": 10.6, "y": 6.0},
        {"id": "Q-B1", "label": "Effectiveness\nrecall + specificity", "kind": "softgoal", "x": 2.6, "y": 1.7},
        {"id": "Q-B2", "label": "Fairness gap <=5pp", "kind": "softgoal", "x": 6.6, "y": 1.7},
        {"id": "BR-B1", "label": "No automated adverse\naction", "kind": "resource", "x": 12.0, "y": 4.0},
        {"id": "BR-B2", "label": "Exclude identifiers +\ndemographics", "kind": "resource", "x": 12.3, "y": 1.7},
    ], [
        ("A-1", "A-2", "depends-on", "normal", 0.0, 0.0, 0.28),
        ("A-2", "A-3", "depends-on", "normal", 0.0, 0.0, -0.28),
        ("A-1", "T-B1", "performs", "normal", 0.0, 0.0, -0.3),
        ("A-3", "BR-B2", "governs audit", "rule", 0.0, 0.35, 0.0, 35, 28),
        ("G-B1", "G-B2", "supports", "normal"),
        ("G-B2", "G-B2a", "AND", "normal", 0.0, -0.2, 0.28),
        ("G-B2", "G-B2b", "AND", "normal", 0.0, 0.2, 0.28),
        ("IND-1", "Q-B1", "", "normal", 0.0, -0.25, -0.1),
        ("IND-2", "Q-B1", "", "normal", 0.0, 0.25, -0.1),
        ("G-B2b", "T-B1", "supports", "normal", 0.0, 0.25, 0.15),
        ("BR-B1", "T-B1", "constrains", "rule", 0.0, 0.55, 0.0, 0, 35),
    ], annotations=[(7, 1.02, "Cross-view trace: IND-1 and IND-2 -> G-A1")])


def draw_gr4ml_analytics_view():
    draw_gr4ml_view("gr4ml_analytics_view.png", "GR4ML Analytics Design View", [
        {"id": "G-A1", "label": "Estimate P(default)\nfor next month", "kind": "analytics_goal", "x": 6.5, "y": 8.4},
        {"id": "T-A1", "label": "Supervised Binary\nClassification", "kind": "task", "x": 6.5, "y": 6.6},
        {"id": "ALG-1", "label": "Logistic Regression", "kind": "algorithm", "x": 3.2, "y": 4.9},
        {"id": "ALG-2", "label": "Random Forest", "kind": "algorithm", "x": 9.8, "y": 4.9},
        {"id": "T-A2", "label": "Holdout Evaluation", "kind": "task", "x": 6.5, "y": 2.8},
        {"id": "IND-A1", "label": "Holdout Confusion\nMatrix & Metrics", "kind": "indicator", "x": 10.8, "y": 2.8},
        {"id": "Q-A1", "label": "Recall/Specificity\nBalance", "kind": "softgoal", "x": 3.0, "y": 1.5},
        {"id": "Q-A2", "label": "Model\nInterpretability", "kind": "softgoal", "x": 6.5, "y": 4.2},
        {"id": "Q-A3", "label": "Reproducibility", "kind": "softgoal", "x": 10.8, "y": 1.5},
        {"id": "BR-A1", "label": "Default = class 1\npositive probability", "kind": "resource", "x": 12.0, "y": 8.2},
    ], [
        ("G-A1", "T-A1", "AND", "normal", 0.0, 0.65, 0.0),
        ("T-A1", "ALG-1", "OR", "normal", 0.0, 0.55, -0.2),
        ("T-A1", "ALG-2", "OR", "normal", 0.0, 0.55, 0.2),
        ("ALG-1", "T-A2", "evaluate", "normal", 0.0, -0.1, -0.15),
        ("ALG-2", "T-A2", "evaluate", "normal", 0.0, 0.1, -0.15),
        ("T-A2", "IND-A1", "produces", "normal", 0.0, 0.0, 0.35),
        ("T-A1", "BR-A1", "enforces", "rule", 0.0, 0.1, 0.3),
        ("ALG-1", "Q-A1", "++", "normal", 0.0, 0.35, 0.1),
        ("ALG-1", "Q-A2", "+", "normal", 0.0, -0.2, -0.45),
        ("ALG-2", "Q-A2", "-", "normal", 0.0, 0.43, -0.45),
        ("T-A2", "Q-A3", "+", "normal", 0.0, 0.35, 0.15),
    ])


def draw_gr4ml_diagrams():
    draw_gr4ml_business_view()
    draw_gr4ml_analytics_view()
    draw_gr4ml_view("gr4ml_dataprep_view.png", "GR4ML Data Preparation View", [
        {"id": "DS-1", "label": "Raw UCI Credit CSV\n[30,000 x 24]", "kind": "data_store", "x": 2.0, "y": 7.4, "width": 2.8, "height": 1.15},
        {"id": "T-D1", "label": "Load & Target Mapping", "kind": "task", "x": 5.4, "y": 7.4},
        {"id": "T-D2", "label": "Remove ID & Protected Attributes", "kind": "task", "x": 9.0, "y": 7.4},
        {"id": "T-D3", "label": "Stratified Split & Train-Only Scaling", "kind": "task", "x": 9.0, "y": 4.7},
        {"id": "DS-2", "label": "Clean Stratified Train/Test Split\n[19 features]", "kind": "data_store", "x": 9.0, "y": 2.0, "width": 3.1, "height": 1.15},
        {"id": "Q-D1", "label": "Schema Completeness", "kind": "softgoal", "x": 2.8, "y": 4.7},
        {"id": "Q-D2", "label": "Zero Test Leakage", "kind": "softgoal", "x": 5.4, "y": 2.0},
        {"id": "BR-D1", "label": "Map default = 1", "kind": "resource", "x": 2.8, "y": 3.3},
        {"id": "BR-D2", "label": "Exclude ID + demographics", "kind": "resource", "x": 5.6, "y": 3.3},
        {"id": "T-A1", "label": "Analytics classification\n(cross-view)", "kind": "task", "x": 12.0, "y": 7.4},
    ], [
        ("DS-1", "T-D1", "input", "normal"),
        ("T-D1", "T-D2", "then", "normal"),
        ("T-D2", "T-D3", "then", "normal"),
        ("T-D3", "DS-2", "produces", "normal"),
        ("DS-2", "T-A1", "analytics trace", "normal", 0.0, 0.2, 0.2),
        ("T-D1", "Q-D1", "+", "normal", 0.05, -0.1, 0.1),
        ("T-D3", "Q-D2", "++", "normal", 0.08, -0.1, 0.15),
        ("T-D1", "BR-D1", "rule", "rule", 0.0, -0.1, 0.0),
        ("T-D2", "BR-D2", "rule", "rule", 0.0, 0.2, 0.0),
    ])


def draw_pattern_diagrams():
    figure, axis = plt.subplots(figsize=(11.5, 6.5))
    figure.patch.set_facecolor("#fbfcfa")
    axis.set_facecolor("#fbfcfa")
    axis.set_xlim(0, 12)
    axis.set_ylim(0, 7)
    axis.axis("off")

    def box(x, y, width, height, title, body, color):
        axis.add_patch(FancyBboxPatch((x - width / 2, y - height / 2), width, height,
                                      boxstyle="round,pad=0.06,rounding_size=0.12",
                                      facecolor=color, edgecolor="#29423a", linewidth=1.4, zorder=2))
        axis.text(x, y + 0.18, title, ha="center", va="center", fontsize=10,
                  weight="bold", color="#18332b", zorder=3)
        axis.text(x, y - 0.2, body, ha="center", va="center", fontsize=8,
                  color="#29423a", zorder=3)

    box(6, 5.9, 4.0, 1.0, "app.py | Presentation Layer", "Streamlit UI, input controls, metrics, safeguards", "#dce9df")
    box(6, 4.15, 4.0, 1.0, "credit_risk/service.py\nApplication Layer", "evaluate(), predict(), threshold workflow", "#e8eef5")
    box(3.0, 2.1, 4.2, 1.0, "credit_risk/strategies.py\nDomain / ML Layer", "ModelStrategy + estimator implementations", "#f4e5c5")
    box(9.0, 2.1, 4.2, 1.0, "credit_risk/data.py\nData Access Layer", "Bundled CSV adapter + stratified split", "#f2e6e6")
    for start, end in [((6, 5.38), (6, 4.68)), ((5.2, 3.65), (3.8, 2.62)), ((6.8, 3.65), (8.2, 2.62))]:
        axis.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=15,
                                       linewidth=1.5, color="#45685a", zorder=1))
    axis.text(6, 0.7, "Dependency direction: Presentation -> Service -> Model Strategy / Data Access",
              ha="center", fontsize=9, color="#344d43")
    axis.set_title("Pattern 1 | Layered Architecture", fontsize=15, weight="bold", color="#18332b", pad=15)
    save_figure(figure, ASSETS / "pattern_layered.png", dpi=240)
    plt.close(figure)

    figure, axis = plt.subplots(figsize=(11.5, 6.5))
    figure.patch.set_facecolor("#fbfcfa")
    axis.set_facecolor("#fbfcfa")
    axis.set_xlim(0, 12)
    axis.set_ylim(0, 7)
    axis.axis("off")
    box(6, 5.9, 4.1, 0.95, "ModelStrategy (Protocol)", "Abstract build() contract", "#e8eef5")
    box(3.1, 3.85, 3.8, 1.0, "LogisticRegressionStrategy", "Scaled, class-weighted estimator", "#f4e5c5")
    box(8.9, 3.85, 3.8, 1.0, "RandomForestStrategy", "Balanced tree ensemble", "#f4e5c5")
    box(3.1, 1.45, 3.8, 0.9, "evaluate()", "Selects strategy, trains, scores metrics", "#dce9df")
    box(8.9, 1.45, 3.8, 0.9, "predict()", "Uses fitted estimator + threshold", "#dce9df")
    for start, end in [((5.35, 5.42), (3.8, 4.36)), ((6.65, 5.42), (8.2, 4.36))]:
        axis.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=15,
                                       linewidth=1.5, color="#45685a", zorder=1))
    for start, end in [((3.1, 1.94), (3.1, 3.32)), ((8.9, 1.94), (8.9, 3.32))]:
        axis.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=15,
                                       linewidth=1.5, color="#45685a", zorder=1))
    axis.text(6, 0.35, "Polymorphic selection: either concrete strategy satisfies the same service contract",
              ha="center", fontsize=9, color="#344d43")
    axis.set_title("Supporting Design Pattern | Strategy", fontsize=15, weight="bold", color="#18332b", pad=15)
    save_figure(figure, ASSETS / "pattern_strategy.png", dpi=240)
    plt.close(figure)

    figure, axis = plt.subplots(figsize=(11.5, 6.5))
    figure.patch.set_facecolor("#fbfcfa")
    axis.set_facecolor("#fbfcfa")
    axis.set_xlim(0, 12)
    axis.set_ylim(0, 7)
    axis.axis("off")

    def service_box(x, y, width, title, body, color):
        axis.add_patch(FancyBboxPatch(
            (x - width / 2, y - 0.52), width, 1.04,
            boxstyle="round,pad=0.06,rounding_size=0.12", facecolor=color,
            edgecolor="#29423a", linewidth=1.4, zorder=2,
        ))
        axis.text(x, y + 0.18, title, ha="center", va="center", fontsize=10,
                  weight="bold", color="#18332b", zorder=3)
        axis.text(x, y - 0.2, body, ha="center", va="center", fontsize=8,
                  color="#29423a", zorder=3)

    service_box(2.1, 5.1, 3.2, "Streamlit client", "human analyst workflow", "#dce9df")
    service_box(2.1, 2.0, 3.2, "External API client", "independent second client", "#e8eef5")
    service_box(6.0, 3.55, 3.4, "FastAPI server", "POST /predict | GET /metrics", "#f4e5c5")
    service_box(10.0, 3.55, 3.4, "Shared application service", "training, scoring, holdout metrics", "#dce9df")
    for start, end, label in [
        ((3.75, 5.0), (4.55, 4.0), "HTTP request / response"),
        ((3.75, 2.1), (4.55, 3.1), "HTTP request / response"),
        ((7.75, 3.55), (8.25, 3.55), "service calls"),
    ]:
        axis.add_patch(FancyArrowPatch(start, end, arrowstyle="<->", mutation_scale=14,
                                       linewidth=1.4, color="#45685a", zorder=1))
        axis.text((start[0] + end[0]) / 2, (start[1] + end[1]) / 2 + 0.3,
                  label, ha="center", fontsize=7.2, color="#344d43",
                  bbox={"facecolor": "#fbfcfa", "edgecolor": "none", "pad": 1}, zorder=3)
    axis.text(6, 0.55, "The API is an optional second client; Streamlit continues to call the same service directly.",
              ha="center", fontsize=8.5, color="#344d43")
    axis.set_title("Pattern 2 | Model-as-a-Service (Client-Server)", fontsize=15,
                   weight="bold", color="#18332b", pad=15)
    save_figure(figure, ASSETS / "pattern_model_service.png", dpi=240)
    plt.close(figure)


def notebook_cells(evidence, test_count):
    markdown = nbformat.v4.new_markdown_cell
    code = nbformat.v4.new_code_cell
    cells = [
        markdown(f"""# AIMLZG546 | Assignment I | Group 54

## Credit Default Risk Review Prototype

**Group members and contributions:**

{member_markdown_table()}

"""),
        markdown("""## 1. Problem formulation and scope

**Domain:** financial risk management for existing credit-card accounts. **Problem statement:** use an account's historical credit limit, repayment status, billed amounts, and payment amounts to estimate the likelihood of payment default next month and prioritize an account for review by a credit-risk analyst. The prototype uses the UCI *Default of Credit Card Clients* dataset: 30,000 Taiwan credit-card accounts and a binary next-month default label. It is an educational benchmark, not a production credit-scoring service.

**Users and stakeholders:** credit-risk analyst (reviews a flag), model owner (validates and monitors), data steward (governs permitted use), customer (affected stakeholder), and service operator (availability/security). The prototype excludes the row identifier and sex, education, marital-status, and age fields from model inputs. Its controls are illustrative median-based profiles, not customer records.

**Dataset citation:** Yeh, I.-C. and Lien, C.-H. (2009), “The comparisons of data mining techniques for the predictive accuracy of probability of default of credit card clients,” *Expert Systems with Applications*, 36(2), 2473-2480. UCI Machine Learning Repository, [Default of Credit Card Clients](https://archive.ics.uci.edu/dataset/350/default+of+credit+card+clients).
"""),
        markdown(f"""## 2. Requirements and measurable goals

### GR4ML Measurable Goals & Indicators Table

{markdown_table(GQI_HEADERS, build_gqi_rows(evidence))}

### Functional and Nonfunctional Requirements

| ID | Requirement | Verification / acceptance |
|---|---|---|
| FR-1 | Load the bundled UCI account data; remove ID and excluded demographic fields; encode next-month default as positive class 1. | Schema, row-count, exclusion, and target checks. |
| FR-2 | Use a reproducible stratified 80/20 train/test split with seed 42. | Repeated split has the same records and class proportions. |
| FR-3 | Allow logistic-regression and random-forest model strategies. | Both strategies train and return valid probabilities. |
| FR-4 | Display recall, specificity, precision, accuracy, and a configurable threshold. | UI metrics match service results; test the threshold boundary. |
| FR-5 | Return estimated default probability and flag accounts at or above the threshold for analyst review. | Positive-class and boundary tests. |
| NFR-1 Effectiveness | Before any pilot, independently validate on later, representative data and meet default recall >=70% and specificity >=70%; otherwise block release. | Temporal/external holdout with a documented threshold and confusion matrix. Current benchmark does not meet both targets at threshold 0.50. |
| NFR-2 Fairness | Before any pilot, measure relevant subgroup error-rate gaps; require recall gap <=5 percentage points or block release. | Lawful, governed audit dataset; demographic fields are audit-only and excluded from model inputs. |
| NFR-3 Privacy | Do not retain identifiers, account inputs, or scores in the prototype; do not transmit data externally. | Inspect code, logging, and storage. |
| NFR-4 Performance / reliability | After model load, scoring p95 <2 seconds at 10 requests/second; production availability >=99% monthly. | Reference-environment load test and monitored SLO report. |

Targets are release criteria, not achieved claims. The notebook's random holdout is only a reproducible coursework demonstration and cannot substitute for temporal validation or governance.
"""),
        markdown("""## 3. GR4ML views

Notation used consistently: **G** = goal, **Q** = quality/softgoal, **A** = actor, **T** = task, **R** = resource, **BR** = business/data rule. AND-decomposition requires all child goals; OR-decomposition denotes alternatives. Dependencies are stated explicitly.

### 3.1 Business View

| Element | GR4ML statement |
|---|---|
| A-1 | Actor: credit-risk analyst reviews context and makes account decisions. |
| A-2 | Actor: ML service estimates next-month default probability. |
| A-3 | Actor: data steward governs provenance and fairness audit. |
| G-B1 | Strategic Goal: support timely review of elevated-risk accounts. |
| G-B2 | Decision Goal: prioritize accounts for analyst review (AND decomposition). |
| G-B2a | Valid feature input. |
| G-B2b | Risk estimate and review flag. |
| IND-1 | Indicator: Default Recall >=70% on independent validation; linked to Q-B1 and G-A1. |
| IND-2 | Indicator: Specificity >=70% on independent validation; linked to Q-B1 and G-A1. |
| Q-B1 | Softgoal: Decision Effectiveness, measured by IND-1 and IND-2. |
| Q-B2 | Softgoal: Demographic Fairness; maximum recall gap <=5pp before pilot. |
| T-B1 | Task: Analyst Account Review using approved conventional evidence. |
| BR-B1 | No automated credit approval, pricing, collections, adverse action, or customer decision. |
| BR-B2 | Exclude direct identifiers and demographics from model inputs; govern audit use. |

**Flow and dependencies:** G-B1 supports G-B2; G-B2 AND-decomposes to valid input (G-B2a) and an estimate/flag (G-B2b). G-B2b supports analyst review T-B1. IND-1 and IND-2 measure Q-B1 and trace to analytics goal G-A1. A-1 retains decision authority; A-3 governs approved data and audit use.

![GR4ML Business View](submission_assets/gr4ml_business_view.png)

### 3.2 Analytics Design View

| Element | GR4ML statement |
|---|---|
| G-A1 | Analytics Goal: estimate next-month P(default) and support review. |
| T-A1 | Analytics Task: Supervised Binary Classification. |
| ALG-1 | Algorithm: Logistic Regression. |
| ALG-2 | Algorithm: Random Forest. |
| T-A2 | Task: evaluate the holdout confusion matrix and metrics at threshold 0.50. |
| IND-A1 | Indicator: Holdout Confusion Matrix & Metrics. |
| Q-A1 | Softgoal: Recall/Specificity Balance. |
| Q-A2 | Softgoal: Model Interpretability. |
| Q-A3 | Softgoal: Reproducibility (fixed seed and split). |
| BR-A1 | Default is positive class 1; select predict_proba column by class label. |

**OR decomposition:** T-A1 selects ALG-1 or ALG-2. Both implement the same model contract. Logistic regression contributes positively to interpretability; the less interpretable forest is a trade-off.

![GR4ML Analytics Design View](submission_assets/gr4ml_analytics_view.png)

### 3.3 Data Preparation View

| Element | GR4ML statement |
|---|---|
| DS-1 | Data Source/Store: Raw UCI Credit CSV [30,000 x 24]. |
| T-D1 | Data Preparation Task: Load & Target Mapping. |
| T-D2 | Data Preparation Task: Remove ID & Protected Attributes. |
| T-D3 | Data Preparation Task: Stratified Split & Train-Only Scaling. |
| DS-2 | Data Source/Store: Clean Stratified Train/Test Split [19 features]. |
| Q-D1 | Softgoal: Schema Completeness. |
| Q-D2 | Softgoal: Zero Test Leakage. |
| BR-D1 | Map default=1 and preserve predictor names/order. |
| BR-D2 | Exclude ID/demographics from training; fit transforms on training data only. |

**Traceability:** G-B1 -> G-B2 -> (AND) G-B2a / G-B2b -> T-B1; IND-1 / IND-2 -> G-A1 -> T-A1 -> (OR) ALG-1 / ALG-2 -> T-A2 -> IND-A1; DS-1 -> T-D1 -> T-D2 -> T-D3 -> DS-2 -> T-A1. Q-D1 and Q-D2 are supported by the data-preparation controls.

![GR4ML Data Preparation View](submission_assets/gr4ml_dataprep_view.png)
"""),
        markdown("""## 4. Top three quality requirements

1. **Effectiveness and decision safety:** false negatives can hide repayment risk, while false positives waste analyst capacity and may unfairly burden customers. Recall and specificity must both be measured. Release gate: each >=70% on independent later-period data, with no automated customer action. At threshold 0.50, the current random holdout does not meet both targets for either strategy; no model is presented as ready for use.
2. **Fairness:** historical lending outcomes may encode structural and policy bias. Sex, education, marital status, and age are excluded from model inputs; this alone does not remove proxy bias. Before a pilot, conduct a lawful, governed subgroup audit and block use if the relevant-group recall gap exceeds 5 percentage points or evidence is inadequate.
3. **Privacy and security:** financial account data is sensitive. The local prototype has no customer storage or external transmission and uses a de-identified public benchmark. A real system would additionally require access control, encryption, audit, retention limits, threat modeling, and regulatory review.
"""),
        markdown(textwrap.dedent("""\
            ## 5. System Architecture and Architectural Patterns

            ![Credit risk system architecture](submission_assets/architecture.png)

            **Non-ML components:** Streamlit UI, FastAPI server, application service, cache/registry, data adapter, and governance logger. **ML components:** training/evaluation, model strategies, fitted artifact, and scoring. Streamlit remains a direct service client; FastAPI is an optional second client. Inputs and predictions are not persisted; there is no automated credit action.

            ### Pattern 1 - Layered Architecture

            ![Layered architecture pattern](submission_assets/pattern_layered.png)

            ### Pattern 2 - Model-as-a-Service (Client-Server)

            ![Model-as-a-Service pattern](submission_assets/pattern_model_service.png)

            ### Supporting Design Pattern - Strategy (GoF)

            ![Strategy pattern](submission_assets/pattern_strategy.png)

            | Pattern Name | Context / Problem | Solution / Structure | Components Mapped in Code | Advantages | Trade-offs |
            |---|---|---|---|---|---|
            | Layered Architecture | UI, ML workflow, and data loading would otherwise be tightly coupled. | Presentation calls service; service coordinates model and data access. | `app.py` -> `credit_risk/service.py` -> `credit_risk/strategies.py` and `credit_risk/data.py`; Streamlit cache stores fitted models. | Clear responsibilities, focused tests, and independent UI/model evolution. | Additional indirection; cache is process-local and must be invalidated after model/data changes. |
            | Model-as-a-Service (Client-Server) | Multiple clients need prediction without coupling to model internals. | FastAPI exposes `POST /predict` and `GET /metrics`; endpoints reuse the application service. | `api.py`, FastAPI, `credit_risk/service.py`; Streamlit remains a direct client. | Shared inference contract; independently testable HTTP interface. | Adds deployment, validation, and service availability responsibilities. |
            | Supporting design pattern: Strategy (GoF) | Compare or replace estimators without duplicating evaluation and inference logic. | Common `ModelStrategy.build()` contract with interchangeable implementations. | `ModelStrategy`, both strategy classes, `get_strategy()`, `fit_model()`, `evaluate_model()`, `predict()`. | Add models without rewriting use cases. | A design pattern, not an architectural pattern; implementations must honor model contracts. |

            Technologies: Python, pandas, scikit-learn Pipelines, Streamlit, matplotlib, and python-docx.
            """)),
        markdown("""## 6. Implementation and reproducibility

    Install runtime dependencies with `python -m pip install -r requirements.txt`; start the UI with `streamlit run app.py`. Start the optional Client-Server API with `uvicorn api:app --reload`. Install verification tools from `requirements-dev.txt`; run tests with `pytest`.

    The bundled CSV keeps runs independent of network access. The Streamlit interface selects a strategy and threshold, displays held-out metrics, and exposes six illustrative account controls; the other 13 model inputs use dataset medians. These controls are not live customer inputs.

    **Safeguard:** this is an educational prototype only, not for credit approval, pricing, collections, adverse action, or other real customer decisions.
"""),
        code("""from credit_risk.data import load_data, split_data
from credit_risk.service import evaluate_model, fit_model, predict
from credit_risk.strategies import STRATEGIES, get_strategy

features, target = load_data()
print(f"Rows: {len(features)} | model features: {features.shape[1]} | defaults: {int(target.sum())}")
print("Split sizes:", [len(part) for part in split_data()])
print("Excluded model inputs: ID, SEX, EDUCATION, MARRIAGE, AGE")"""),
        markdown("""## 7. Holdout metrics and threshold sweep

    The following sweep computes recall, specificity, and precision at thresholds 0.30 to 0.60 in 0.05 steps for both strategies. Both >=70% targets must hold at the same threshold to pass.
    """),
        code("""results = {}
for name in STRATEGIES:
    model = fit_model(get_strategy(name))
    metrics = evaluate_model(model)
    results[name] = metrics
    print(name, {metric: round(float(metrics[metric]), 4) for metric in
                 ("recall", "specificity", "precision", "accuracy")})"""),
        code("""import pandas as pd

threshold_rows = []
for name in STRATEGIES:
    model = fit_model(get_strategy(name))
    for threshold in [round(0.30 + 0.05 * i, 2) for i in range(7)]:
        metrics = evaluate_model(model, threshold)
        threshold_rows.append({
            "strategy": name,
            "threshold": threshold,
            "recall": metrics["recall"],
            "specificity": metrics["specificity"],
            "precision": metrics["precision"],
        })
sweep = pd.DataFrame(threshold_rows)
display(sweep.style.format({"threshold": "{:.2f}", "recall": "{:.2%}",
                            "specificity": "{:.2%}", "precision": "{:.2%}"}))
both_targets = sweep[(sweep["recall"] >= 0.70) & (sweep["specificity"] >= 0.70)]
print("A tested threshold meets both >=70% targets:", not both_targets.empty)
if both_targets.empty:
    print("No tested strategy/threshold pair meets both targets on this fixed holdout.")
else:
    print(both_targets[["strategy", "threshold"]].to_string(index=False))"""),
        markdown("""## 8. Audit-only subgroup recall

SEX and AGE are loaded separately for the same held-out row indices; they are not model features. The maximum subgroup recall gap is compared with the <=5 percentage-point goal. This is a coursework audit, not a governed fairness assessment.
"""),
        code("""from credit_risk.fairness import audit_recall_by_group

fairness_results = {}
for name in STRATEGIES:
    audit = audit_recall_by_group(fit_model(get_strategy(name)))
    fairness_results[name] = audit
    print(name, audit["recall_by_group"])
    print(f"Maximum recall gap: {audit['max_gap_pp']:.2f} percentage points; "
          f"{'MEETS' if audit['meets_target'] else 'DOES NOT MEET'} <=5pp target")"""),
        markdown("""## 9. Prediction and automated verification

The test cell exercises the same data, split, target, probability, threshold, strategy, fairness-input, and API conditions documented in the report.
"""),
        code("""model = fit_model(get_strategy("Logistic regression"))
example = features.median().to_frame().T
result = predict(model, example, threshold=0.50)
print({key: round(value, 4) if isinstance(value, float) else value
       for key, value in result.items()})
print("Educational example only; not a customer-level financial decision.")"""),
    code(f"""import pytest

exit_code = pytest.main(["-q", "tests/test_credit_risk.py"])
assert int(exit_code) == 0
print("{test_count} tests passed")"""),
    ]
    cells.append(markdown("## 10. Appendix - Complete Source Code\n\nThe complete implementation and tests are reproduced below so the notebook can be reviewed alongside the submission bundle."))
    for source_path in SOURCE_FILES:
        source = source_path.read_text(encoding="utf-8").rstrip()
        relative_path = source_path.relative_to(ROOT).as_posix()
        cells.append(markdown(f"### `{relative_path}`\n\n```python\n{source}\n```"))
    return cells


def build_notebook(evidence, test_count):
    notebook = nbformat.v4.new_notebook(cells=notebook_cells(evidence, test_count))
    for index, cell in enumerate(notebook.cells, start=1):
        cell_id = f"group54-cell-{index:03d}-{uuid4().hex[:8]}"
        cell.id = cell_id
        cell.metadata["id"] = cell_id
        cell.metadata["language"] = "python" if cell.cell_type == "code" else "markdown"
    notebook.metadata.kernelspec = {
        "display_name": "Python 3",
        "language": "python",
        "name": "python3",
    }
    notebook.metadata.language_info = {"name": "python", "version": "3"}
    diagram_names = [
        "architecture.png", "gr4ml_business_view.png", "gr4ml_analytics_view.png",
        "gr4ml_dataprep_view.png", "pattern_layered.png", "pattern_strategy.png",
        "pattern_model_service.png",
    ]
    for cell in notebook.cells:
        if cell.cell_type != "markdown":
            continue
        for filename in diagram_names:
            marker = f"submission_assets/{filename}"
            if marker in cell.source:
                encoded = base64.b64encode((ASSETS / filename).read_bytes()).decode("ascii")
                cell.source = cell.source.replace(marker, f"data:image/png;base64,{encoded}")
    nbformat.write(notebook, ROOT / "54.ipynb")


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    return table


def build_report(evidence, test_count, output_path=None):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor(37, 55, 48)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor(24, 51, 43)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AIMLZG546\nSoftware Engineering for Machine Learning\nAssignment I")
    run.bold = True
    run.font.size = Pt(20)
    subtitle = document.add_paragraph("Group 54 | Credit Default Risk Review Prototype")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Assignment report and implementation evidence")

    document.add_heading("Group Details", level=1)
    document.add_paragraph("Group No: 54. Member contributions total 100%.")
    member_table = add_table(document, MEMBER_HEADERS, MEMBER_ROWS)
    member_table.autofit = False
    member_widths = [0.45, 0.85, 1.0, 4.2, 0.75]
    for row in member_table.rows:
        for cell, width in zip(row.cells, member_widths):
            cell.width = Inches(width)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)

    document.add_heading("1. Problem Formulation", level=1)
    document.add_paragraph(
        "Domain: financial risk management for existing credit-card accounts. Using account "
        "repayment status, credit limit, billed amounts and payment amounts, the prototype "
        "estimates the likelihood of payment default next month and flags an account for "
        "credit-risk analyst review. It uses the UCI Default of Credit Card Clients dataset: "
        "30,000 historical Taiwan accounts and a binary next-month default outcome."
    )
    document.add_paragraph(
        "The analyst retains all account-level decision authority. This educational prototype "
        "must not be used for credit approval, pricing, collections, adverse action, or other "
        "customer-level decisions. The bundled CSV has no source ID in model inputs; sex, "
        "education, marital status, and age are excluded from model inputs. UI controls are "
        "illustrative median-based profiles, not customer data."
    )
    document.add_paragraph(
        "Users and stakeholders: the credit-risk analyst reviews flags and retains decision authority; "
        "the model owner validates and monitors; the data steward governs permitted use; customers "
        "are affected stakeholders; and the service operator is responsible for availability and security."
    )
    document.add_paragraph(
        "Dataset source: Yeh, I.-C. and Lien, C.-H. (2009), Expert Systems with Applications, "
        "36(2), 2473-2480; UCI Machine Learning Repository, Default of Credit Card Clients "
        "(https://archive.ics.uci.edu/dataset/350/default+of+credit+card+clients)."
    )

    document.add_heading("2. Requirements and Measurable Goals", level=1)
    document.add_heading("GR4ML Measurable Goals & Indicators Table", level=2)
    gqi_table = add_table(document, GQI_HEADERS, build_gqi_rows(evidence))
    gqi_widths = [0.55, 0.95, 1.35, 1.15, 1.25, 1.85]
    for row in gqi_table.rows:
        for cell, width in zip(row.cells, gqi_widths):
            cell.width = Inches(width)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)
    document.add_heading("Functional and Nonfunctional Requirements", level=2)
    add_table(document, ["ID", "Specification", "Acceptance / verification"], [
        ["FR-1", "Load bundled UCI data, remove ID and excluded demographics, map default to class 1.", "Schema, exclusion, and target checks."],
        ["FR-2", "Use stratified 80/20 train/test split, random_state=42.", "Repeated split and class proportions are reproducible."],
        ["FR-3", "Select logistic regression or random forest.", "Both strategies fit and return valid probabilities."],
        ["FR-4", "Display recall, specificity, precision, accuracy; adjust threshold.", "UI metrics match service; threshold-boundary check."],
        ["FR-5", "Return default probability and flag at/above threshold for analyst review.", "Positive-class and decision-boundary tests."],
        ["NFR-1 Effectiveness", "Independent later-period data must show recall >=70% and specificity >=70% before a pilot.", "Temporal/external validation; block release if unmet."],
        ["NFR-2 Fairness", "Relevant-group recall gap <=5 percentage points before a pilot.", "Lawfully governed subgroup audit; demographics audit-only."],
        ["NFR-3 Privacy", "No identifier/input/score retention or external transmission in prototype.", "Code, log, and storage inspection."],
        ["NFR-4 Performance", "Scoring p95 <2 s at 10 req/s after model load; availability >=99% monthly in production.", "Reference load test and SLO report."],
    ])
    document.add_paragraph(
        "These are proposed acceptance targets, not achieved claims. The threshold sweep results "
        + ("include a threshold meeting both targets on this fixed holdout; this does not establish independent performance."
           if evidence["threshold_target_met"] else "show that neither strategy meets both targets between 0.30 and 0.60 on this fixed holdout.")
    )

    document.add_heading("3. GR4ML Views", level=1)
    document.add_paragraph(
        "Notation: G = goal, Q = quality/softgoal, A = actor, T = task, R = resource, "
        "BR = business/data rule. AND-decomposition requires all child goals; OR-decomposition "
        "denotes alternatives; dependencies state who relies on whom."
    )
    document.add_heading("3.1 Business View", level=2)
    add_table(document, ["Element", "GR4ML description"], [
        ["A-1 / A-2 / A-3", "Actors: credit-risk analyst / ML service / data steward."],
        ["G-B1", "Strategic Goal: support timely review of elevated-risk accounts; supports G-B2."],
        ["G-B2", "Decision Goal: prioritize accounts for analyst review; AND-decomposes to G-B2a and G-B2b."],
        ["G-B2a / G-B2b", "Valid feature input / risk estimate and review flag."],
        ["IND-1 / IND-2", "Indicators: Default Recall >=70%; Specificity >=70%; both linked to Q-B1 and G-A1."],
        ["Q-B1 / Q-B2", "Softgoals: Decision Effectiveness; Demographic Fairness <=5pp recall gap."],
        ["T-B1", "Task: Analyst Account Review using approved conventional evidence."],
        ["BR-B1", "No automated credit approval, pricing, collections, adverse action, or customer decision."],
        ["BR-B2", "Exclude direct identifiers and demographics from model inputs; govern audit use."],
        ["Flow / dependencies", "G-B1 supports G-B2; G-B2 AND-decomposes to G-B2a/G-B2b. IND-1/IND-2 measure Q-B1 and trace to G-A1; G-B2b supports T-B1."],
    ])
    document.add_picture(str(ASSETS / "gr4ml_business_view.png"), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("3.2 Analytics Design View", level=2)
    add_table(document, ["Element", "GR4ML description"], [
        ["G-A1", "Analytics Goal: estimate next-month P(default) and support review."],
        ["T-A1", "Analytics Task: Supervised Binary Classification."],
        ["ALG-1 / ALG-2", "Algorithms: Logistic Regression OR Random Forest."],
        ["T-A2", "Evaluate holdout confusion matrix and metrics at threshold 0.50."],
        ["IND-A1", "Indicator: Holdout Confusion Matrix & Metrics."],
        ["Q-A1 / Q-A2 / Q-A3", "Softgoals: Recall/Specificity Balance, Model Interpretability, Reproducibility."],
        ["BR-A1", "Default is class 1; select predict_proba column by class label."],
    ])
    document.add_picture(str(ASSETS / "gr4ml_analytics_view.png"), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("3.3 Data Preparation View", level=2)
    add_table(document, ["Element", "GR4ML description"], [
        ["DS-1", "Data Source/Store: Raw UCI Credit CSV [30,000 x 24]."],
        ["T-D1", "Data Preparation Task: Load & Target Mapping."],
        ["T-D2", "Data Preparation Task: Remove ID & Protected Attributes."],
        ["T-D3", "Data Preparation Task: Stratified Split & Train-Only Scaling."],
        ["DS-2", "Data Source/Store: Clean Stratified Train/Test Split [19 features]."],
        ["Q-D1 / Q-D2", "Softgoals: Schema Completeness and Zero Test Leakage."],
        ["BR-D1 / BR-D2", "Map default=1; preserve schema; exclude ID/demographics; fit transforms on train only."],
        ["Traceability", "G-B1 supports G-B2 -> (AND) G-B2a/G-B2b -> T-B1; IND-1/IND-2 -> G-A1 -> T-A1 -> (OR) ALG-1/ALG-2 -> T-A2 -> IND-A1; DS-1 -> T-D1 -> T-D2 -> T-D3 -> DS-2 -> T-A1."],
    ])
    document.add_picture(str(ASSETS / "gr4ml_dataprep_view.png"), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("4. Top Three Quality Requirements", level=1)
    document.add_paragraph(
          "1. Effectiveness and decision safety: false negatives hide repayment risk; false positives "
          "consume analyst capacity and can burden customers. Measure both recall and specificity; "
          "require each >=70% on independent later-period data. The 0.30-0.60 coursework threshold "
          + ("sweep found at least one pair meeting both metrics on this fixed holdout, not independent evidence."
              if evidence["threshold_target_met"] else "sweep found no pair meeting both metrics on this fixed holdout.")
          + " No release or automated-action claim is made."
    )
    document.add_paragraph(
        "2. Fairness: historic outcomes may reflect structural and policy bias. Excluding demographic "
        "fields does not remove proxy effects. Require a lawful subgroup audit and <=5 percentage-point "
        "relevant-group recall gap before any pilot; block if evidence is inadequate."
    )
    document.add_paragraph(
        "3. Privacy and security: financial account data is sensitive. The local prototype retains no "
        "customer input and makes no external transmission. Real deployment requires access control, "
        "encryption, audit, retention limits, threat modeling, and regulatory review."
    )

    document.add_heading("5. System Architecture", level=1)
    document.add_picture(str(ASSETS / "architecture.png"), width=Inches(6.8))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "Non-ML components: Streamlit UI, FastAPI server, application service, cache/registry, data "
        "adapter, and governance logger. ML components: training/evaluation, model strategies, fitted "
        "artifact, and scoring. Streamlit remains a direct service client; FastAPI is an optional second "
        "client/server path. Inputs and predictions are not persisted; no automated credit action occurs."
    )

    document.add_heading("6. Architectural Patterns and Implementation", level=1)
    document.add_heading("Pattern 1: Layered Architecture", level=2)
    document.add_paragraph(
        "Implemented as app.py presentation -> credit_risk/service.py application workflow -> "
        "credit_risk/strategies.py model abstraction and credit_risk/data.py dataset access. "
        "The service separates fit_model from evaluate_model; Streamlit cache_resource retains the "
        "fitted model across threshold and input changes. Evaluation logs aggregate metrics only."
    )
    document.add_picture(str(ASSETS / "pattern_layered.png"), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Pattern 2: Model-as-a-Service (Client-Server)", level=2)
    document.add_paragraph(
        "api.py exposes POST /predict and GET /metrics through FastAPI and reuses credit_risk.service. "
        "The API is an optional second client; the Streamlit app continues calling the shared service "
        "directly, so both clients use the same model and evaluation behavior."
    )
    document.add_picture(str(ASSETS / "pattern_model_service.png"), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Supporting Design Pattern: Strategy", level=2)
    document.add_paragraph(
        "Within the model layer, ModelStrategy is implemented by LogisticRegressionStrategy and "
        "RandomForestStrategy, registered in STRATEGIES and selected by get_strategy. This is a "
        "GoF design pattern, not an architectural pattern."
    )
    document.add_picture(str(ASSETS / "pattern_strategy.png"), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    pattern_table = add_table(document,
        ["Pattern Name", "Context / Problem", "Solution / Structure", "Components Mapped in Code", "Advantages", "Trade-offs"], [
            ["Layered Architecture", "UI, ML workflow, and data access would otherwise be coupled.",
             "Presentation calls service; service coordinates model and data layers.",
             "app.py -> credit_risk/service.py -> strategies.py and data.py.",
             "Clear responsibilities; focused tests; independent UI/model changes.",
             "Added indirection; cache is process-local and must be invalidated after model/data changes."],
            ["Model-as-a-Service (Client-Server)", "Provide a reusable prediction service to multiple clients without coupling clients to model internals.",
             "FastAPI exposes POST /predict and GET /metrics; endpoints delegate to the shared credit_risk service.",
             "api.py; FastAPI; credit_risk/service.py; Streamlit remains a direct second service client.",
             "Client/server boundary; shared inference behavior; independently testable HTTP contract.",
             "Adds deployment, API validation, and service-availability responsibilities."],
            ["Supporting design pattern: Strategy (GoF)", "Compare or replace estimators without duplicating evaluation/inference logic.",
             "Common build() contract with interchangeable model implementations.",
             "ModelStrategy; LogisticRegressionStrategy; RandomForestStrategy; get_strategy(); evaluate_model(); predict().",
             "Add models without rewriting use cases; consistent comparison workflow.",
             "Not itself an architectural pattern; implementations must honor feature/probability contracts."],
        ])
    for row in pattern_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)

    document.add_heading("7. Implementation Results", level=1)
    document.add_paragraph(
        f"The fixed stratified holdout contains 6,000 records. The automated suite reports {test_count} tests passed. "
        "The tables below are computed at report-generation time from the same bundled CSV, split, "
        "fitted strategies, and threshold sweep used by the notebook code."
    )
    document.add_paragraph(f"Verification result: {test_count} tests passed.")
    baseline_rows = [
        [name, *(f"{metrics[key]:.2%}" for key in ("recall", "specificity", "precision", "accuracy"))]
        for name, metrics in evidence["baseline"].items()
    ]
    add_table(document, ["Strategy", "Default recall", "Specificity", "Precision", "Accuracy"], baseline_rows)
    document.add_paragraph(
        "Results are at threshold 0.50 on this fixed 6,000-row historical holdout. They do not "
        "establish future performance, calibration, fairness, or suitability for customer decisions."
    )
    document.add_heading("Threshold sweep (fixed holdout)", level=2)
    sweep_rows = [
        [row["strategy"], f"{row['threshold']:.2f}", f"{row['recall']:.2%}",
         f"{row['specificity']:.2%}", f"{row['precision']:.2%}"]
        for row in evidence["sweep"]
    ]
    sweep_table = add_table(document, ["Strategy", "Threshold", "Recall", "Specificity", "Precision"], sweep_rows)
    for row in sweep_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)
    document.add_paragraph(
        "Threshold range 0.30 to 0.60 in 0.05 steps: "
        + ("at least one tested strategy/threshold met both >=70% recall and >=70% specificity on this holdout."
           if evidence["threshold_target_met"] else "no tested strategy/threshold met both >=70% recall and >=70% specificity.")
        + " This is not independent validation."
    )
    document.add_heading("Audit-only subgroup recall", level=2)
    fairness_rows = [
        [row["strategy"], row["dimension"], row["group"], f"{row['recall']:.2%}"]
        for row in evidence["fairness_rows"]
    ]
    add_table(document, ["Strategy", "Audit dimension", "Group", "Recall"], fairness_rows)
    for name, audit in evidence["fairness"].items():
        document.add_paragraph(
            f"{name}: maximum recall gap {audit['max_gap_pp']:.2f} percentage points; "
            f"{'meets' if audit['meets_target'] else 'does not meet'} the <=5pp target."
        )
    document.add_paragraph(
        "Coursework audit only, not a legally governed or production fairness assessment. Original SEX and AGE "
        "columns were read separately for the held-out row indices and were never model inputs."
    )

    document.add_heading("8. Application Screenshots", level=1)
    screenshot_items = [
        ("app_logistic_050.png", "Logistic regression at threshold 0.50 with the four held-out evaluation metrics."),
        ("app_forest_changed_threshold.png", "Random forest selected with the review threshold changed from its default."),
        ("app_flagged.png", "Illustrative account controls produce a flagged result for analyst review."),
        ("app_not_flagged.png", "Illustrative account controls produce a result below the selected review threshold."),
    ]
    missing_screenshots = []
    for filename, caption in screenshot_items:
        image_path = ASSETS / filename
        if image_path.exists():
            document.add_paragraph(caption)
            document.add_picture(str(image_path), width=Inches(6.8))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            missing_screenshots.append(filename)
    if missing_screenshots:
        document.add_paragraph("Screenshot capture is pending for: " + ", ".join(missing_screenshots))

    document.add_heading("9. Run Instructions, Limitations and Conclusion", level=1)
    document.add_paragraph(
        "Install dependencies using `python -m pip install -r requirements.txt`; run with "
        "`streamlit run app.py`. The bundled CSV permits offline application use."
    )
    document.add_paragraph("Optional API: `uvicorn api:app --reload`; test suite: `python -m pytest -q`.")
    document.add_paragraph(
        "Limitations: historical Taiwan accounts from 2005; one random split; imbalanced labels; "
        "no temporal/external validation, calibration, confidence intervals, governed subgroup audit, "
        "monitoring, authentication, persistence, load test, or production security review. "
        "Demographic exclusion and class weighting do not guarantee fairness. Never use these "
        "outputs for decisions about real customers."
    )
    document.add_paragraph(
        "The implementation demonstrates requirements traceability through all three GR4ML views, "
        "a layered architecture, and two implemented architectural patterns. Any real-world pilot "
        "would require independent evidence, legal/compliance review, fairness and security controls, "
        "and accountable human governance."
    )

    document.add_heading("10. Appendix - Complete Source Code", level=1)
    document.add_paragraph(
        "The complete source files used by this implementation are reproduced here for review."
    )
    for source_path in SOURCE_FILES:
        relative_path = source_path.relative_to(ROOT).as_posix()
        document.add_heading(relative_path, level=2)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = False
        source_lines = source_path.read_text(encoding="utf-8").replace("\r\n", "\n").rstrip().split("\n")
        for line in source_lines:
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            run.add_break()
    document.save(output_path or ROOT / "54.docx")


def build_submission_bundle():
    archive_path = ROOT / "54_submission.zip"
    required_assets = [
        "credit_default.csv", "architecture.png", "gr4ml_business_view.png",
        "gr4ml_analytics_view.png", "gr4ml_dataprep_view.png", "pattern_layered.png",
        "pattern_model_service.png", "pattern_strategy.png", "app_logistic_050.png",
        "app_forest_changed_threshold.png", "app_flagged.png", "app_not_flagged.png",
    ]
    bundle_files = [
        ROOT / "54.docx", ROOT / "54.ipynb", ROOT / "app.py", ROOT / "api.py",
        ROOT / "build_submission.py", ROOT / "requirements.txt", ROOT / "requirements-dev.txt",
        ROOT / "README.md",
    ]
    bundle_files.extend(sorted((ROOT / "credit_risk").glob("*.py")))
    bundle_files.extend(sorted((ROOT / "tests").glob("*.py")))
    bundle_files.extend(ASSETS / filename for filename in required_assets)
    missing = [path for path in bundle_files if not path.is_file()]
    if missing:
        raise FileNotFoundError("Cannot create submission bundle; missing: " + ", ".join(map(str, missing)))
    with ZipFile(archive_path, "w", ZIP_DEFLATED) as archive:
        for path in bundle_files:
            archive.write(path, path.relative_to(ROOT).as_posix())
    return archive_path


if __name__ == "__main__":
    draw_architecture()
    draw_gr4ml_diagrams()
    draw_pattern_diagrams()
    evidence = collect_submission_evidence()
    test_count, _ = run_test_suite()
    build_notebook(evidence, test_count)
    build_report(evidence, test_count)
    bundle_path = build_submission_bundle()
    print(f"Generated 54.ipynb, 54.docx, diagrams, and {bundle_path.name}; {test_count} tests passed")
